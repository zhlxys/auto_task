import requests
from bs4 import BeautifulSoup
import os
from datetime import datetime
import docx
import re

def search_audit_info():
    search_terms = ["审计资料", "audit information", "审计报告", "audit report"]
    results = []
    
    for term in search_terms:
        try:
            # 使用Bing搜索
            url = f"https://www.bing.com/search?q={term}"
            headers = {
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36"
            }
            response = requests.get(url, headers=headers, timeout=10)
            soup = BeautifulSoup(response.text, 'html.parser')
            
            # 提取搜索结果
            for item in soup.select('li.b_algo')[:5]:  # 取前5个结果
                title = item.select_one('h2 a').text if item.select_one('h2 a') else ""
                link = item.select_one('h2 a')['href'] if item.select_one('h2 a') else ""
                snippet = item.select_one('p').text if item.select_one('p') else ""
                
                if title and link:
                    results.append({
                        "title": title,
                        "link": link,
                        "snippet": snippet,
                        "search_term": term
                    })
        except Exception as e:
            print(f"搜索 {term} 时出错: {e}")
    
    return results

def fetch_content(url):
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36"
        }
        response = requests.get(url, headers=headers, timeout=15)
        response.encoding = response.apparent_encoding
        soup = BeautifulSoup(response.text, 'html.parser')
        
        # 提取正文内容
        content = ""
        for para in soup.find_all(['p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6']):
            content += para.get_text() + "\n"
        
        # 清理内容
        content = re.sub(r'\s+', ' ', content).strip()
        return content[:3000]  # 限制内容长度
    except Exception as e:
        print(f"获取 {url} 内容时出错: {e}")
        return ""

def generate_doc(results, output_path):
    doc = docx.Document()
    doc.add_heading('审计相关资料汇总', 0)
    doc.add_paragraph(f'生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    doc.add_paragraph()
    
    for i, result in enumerate(results, 1):
        doc.add_heading(f'资料 {i}', level=1)
        doc.add_paragraph(f'标题: {result["title"]}')
        doc.add_paragraph(f'链接: {result["link"]}')
        doc.add_paragraph(f'搜索词: {result["search_term"]}')
        doc.add_paragraph(f'摘要: {result["snippet"]}')
        
        # 获取详细内容
        content = fetch_content(result["link"])
        if content:
            doc.add_heading('详细内容', level=2)
            doc.add_paragraph(content)
        
        doc.add_page_break()
    
    doc.save(output_path)

def main():
    # 搜索审计相关信息
    print("正在搜索审计相关资料...")
    results = search_audit_info()
    
    if not results:
        print("未找到相关资料")
        return
    
    # 生成文件名
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = f"audit_data_{timestamp}.docx"
    output_path = os.path.join(os.path.dirname(__file__), output_file)
    
    # 生成文档
    print(f"正在生成文档: {output_file}")
    generate_doc(results, output_path)
    
    print(f"任务完成！文档已保存至: {output_path}")

if __name__ == "__main__":
    main()
