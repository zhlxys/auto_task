#!/usr/bin/env python3
import os
import requests
from bs4 import BeautifulSoup
from datetime import datetime
from docx import Document
from docx.shared import Pt
import re


def sanitize_filename(filename):
    filename = re.sub(r'[<>:"/\\|?*]', '', filename)
    filename = filename.strip()
    return filename[:100]


def create_doc_file(content, title, source, save_dir='.'):
    doc = Document()
    
    title_para = doc.add_heading(title, level=1)
    title_para.runs[0].font.size = Pt(16)
    
    doc.add_heading('来源信息', level=2)
    doc.add_paragraph(f'来源: {source}')
    doc.add_paragraph(f'获取时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    
    doc.add_heading('正文内容', level=2)
    if isinstance(content, list):
        for paragraph in content:
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
    else:
        doc.add_paragraph(str(content))
    
    date_str = datetime.now().strftime("%Y%m%d")
    safe_title = sanitize_filename(title)
    filename = f"{date_str}_{source}_{safe_title}.doc"
    filepath = os.path.join(save_dir, filename)
    
    doc.save(filepath)
    print(f"文档已保存: {filepath}")
    return filepath


def scrape_mof():
    print("开始抓取财政部官网...")
    docs = []
    try:
        url = "http://www.mof.gov.cn/zhengwuxinxi/index.htm"
        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
        }
        response = requests.get(url, headers=headers, timeout=30)
        response.encoding = 'utf-8'
        
        soup = BeautifulSoup(response.text, 'lxml')
        
        items = soup.find_all('a', href=True)
        for item in items[:10]:
            link_text = item.get_text(strip=True)
            link_href = item.get('href', '')
            
            if any(keyword in link_text for keyword in ['审计', '会计', '准则', '制度', '规范']):
                if not link_href.startswith('http'):
                    if link_href.startswith('/'):
                        link_href = 'http://www.mof.gov.cn' + link_href
                    else:
                        link_href = 'http://www.mof.gov.cn/zhengwuxinxi/' + link_href
                
                try:
                    detail_response = requests.get(link_href, headers=headers, timeout=30)
                    detail_response.encoding = 'utf-8'
                    detail_soup = BeautifulSoup(detail_response.text, 'lxml')
                    
                    content_paras = detail_soup.find_all('p')
                    content_texts = [p.get_text(strip=True) for p in content_paras if p.get_text(strip=True)]
                    
                    if content_texts:
                        docs.append({
                            'title': link_text or f"财政部审计相关文件-{len(docs)+1}",
                            'content': content_texts,
                            'source': '财政部'
                        })
                except Exception as e:
                    print(f"抓取详情页面失败: {e}")
                    continue
                    
    except Exception as e:
        print(f"抓取财政部官网出错: {e}")
    
    print(f"从财政部获取到 {len(docs)} 条资料")
    return docs


def scrape_cicpa():
    print("开始抓取中注协官网...")
    docs = []
    try:
        url = "http://www.cicpa.org.cn/news/"
        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
        }
        response = requests.get(url, headers=headers, timeout=30)
        response.encoding = 'utf-8'
        
        soup = BeautifulSoup(response.text, 'lxml')
        
        items = soup.find_all('a', href=True)
        for item in items[:10]:
            link_text = item.get_text(strip=True)
            link_href = item.get('href', '')
            
            if any(keyword in link_text for keyword in ['审计', '准则', '指引', '规定']):
                if not link_href.startswith('http'):
                    if link_href.startswith('/'):
                        link_href = 'http://www.cicpa.org.cn' + link_href
                    else:
                        link_href = 'http://www.cicpa.org.cn/news/' + link_href
                
                try:
                    detail_response = requests.get(link_href, headers=headers, timeout=30)
                    detail_response.encoding = 'utf-8'
                    detail_soup = BeautifulSoup(detail_response.text, 'lxml')
                    
                    content_paras = detail_soup.find_all('p')
                    content_texts = [p.get_text(strip=True) for p in content_paras if p.get_text(strip=True)]
                    
                    if content_texts:
                        docs.append({
                            'title': link_text or f"中注协审计相关文件-{len(docs)+1}",
                            'content': content_texts,
                            'source': '中注协'
                        })
                except Exception as e:
                    print(f"抓取详情页面失败: {e}")
                    continue
                    
    except Exception as e:
        print(f"抓取中注协官网出错: {e}")
    
    print(f"从中注协获取到 {len(docs)} 条资料")
    return docs


def main():
    print("=" * 60)
    print(f"审计资料抓取程序启动 - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 60)
    
    all_docs = []
    
    all_docs.extend(scrape_mof())
    all_docs.extend(scrape_cicpa())
    
    if all_docs:
        print(f"\n共获取到 {len(all_docs)} 条资料，开始生成文档...")
        
        save_dir = '/workspace'
        
        for doc_data in all_docs:
            try:
                create_doc_file(
                    content=doc_data['content'],
                    title=doc_data['title'],
                    source=doc_data['source'],
                    save_dir=save_dir
                )
            except Exception as e:
                print(f"生成文档失败: {e}")
        
        print("\n文档生成完成！")
    else:
        print("\n未获取到任何资料")
    
    print("=" * 60)


if __name__ == "__main__":
    main()
