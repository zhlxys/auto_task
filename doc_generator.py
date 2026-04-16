from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import config
import re


def sanitize_filename(filename):
    return re.sub(r'[<>:"/\\|?*]', '_', filename)


def generate_doc(data):
    doc = Document()
    
    title = doc.add_heading(data['title'], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    info_para = doc.add_paragraph()
    info_para.add_run(f'来源: {data["source"]}\n').bold = True
    info_para.add_run(f'日期: {data["date"]}')
    
    doc.add_paragraph()
    
    paragraphs = data['content'].split('\n')
    for para_text in paragraphs:
        if para_text.strip():
            para = doc.add_paragraph(para_text)
            para.paragraph_format.line_spacing = 1.5
            para.paragraph_format.space_after = Pt(6)
    
    safe_title = sanitize_filename(data['title'])
    if len(safe_title) > 50:
        safe_title = safe_title[:50]
    
    filename = f'{data["date"]}_{data["source"]}_{safe_title}.doc'
    filepath = os.path.join(config.OUTPUT_DIR, filename)
    
    doc.save(filepath)
    return filepath


def generate_docs(data_list):
    filepaths = []
    for data in data_list:
        try:
            filepath = generate_doc(data)
            filepaths.append(filepath)
            print(f'文档已生成: {filepath}')
        except Exception as e:
            print(f'生成文档时出错: {e}')
    return filepaths


if __name__ == '__main__':
    sample_data = {
        'title': '测试文档',
        'content': '这是测试内容。\n\n这是第二段测试内容。',
        'source': '测试来源',
        'date': '20260416'
    }
    generate_doc(sample_data)
